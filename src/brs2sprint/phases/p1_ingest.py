"""Phase 1 — Document ingestion and map-reduce summarization.

Parsers are optional dependencies, resolved at call time:
  .pdf  -> pypdf
  .docx -> python-docx
  .txt/.md -> stdlib
"""

from __future__ import annotations

import re
from pathlib import Path

from ..llm import LLMClient, LLMRequest
from ..prompts import SYSTEM_ANALYST, SUMMARIZE_CHUNK, REDUCE_SUMMARY
from ..schemas import BRSSummary, DocumentChunk

_BULLET_START = re.compile(r"^\s*(?:[-*\u2022]|\d+[.)])\s+")


def _iter_block_items(doc):
    """Yield paragraphs and tables in the order they actually appear in the
    document body.

    python-docx's own `doc.paragraphs` and `doc.tables` are two separate
    flat lists with no ordering between them, so naively concatenating them
    (paragraphs first, then every table) silently detaches each table from
    the heading it belongs to — a requirements table under "4. Business
    Requirements" ends up appended after every other section instead. This
    walks the underlying XML body directly to preserve real document order.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _is_list_paragraph(paragraph) -> bool:
    """True if a Word paragraph carries bullet/numbered list formatting.

    Word represents a list item one of two ways, and real documents use
    both depending on how they were authored: direct `numPr` formatting on
    the paragraph (the common case when typed directly in Word), or a named
    paragraph style such as "List Bullet" / "List Number" that carries the
    numbering at the style level instead. Either way, `.text` itself never
    contains the bullet glyph, so both must be checked.
    """
    pf = paragraph._p.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"
        "/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr"
    )
    if pf is not None:
        return True
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    return "list bullet" in style_name or "list number" in style_name or style_name == "list paragraph"


# --------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------

def load_document(path: str | Path) -> str:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    ext = p.suffix.lower()

    if ext in (".txt", ".md", ".markdown"):
        return p.read_text(encoding="utf-8", errors="replace")

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ImportError("PDF input needs `pip install pypdf`") from exc
        reader = PdfReader(str(p))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise ImportError("DOCX input needs `pip install python-docx`") from exc
        d = docx.Document(str(p))
        blocks = []
        for item in _iter_block_items(d):
            if isinstance(item, docx.text.paragraph.Paragraph):
                text = item.text.strip()
                if text:
                    # Word list items (bulleted/numbered) don't store the glyph
                    # in .text, only in paragraph formatting — restore a marker
                    # so downstream bullet-detection (mock heuristics, and any
                    # regex-based extraction) can still recognise them as list
                    # items rather than plain prose.
                    if _is_list_paragraph(item) and not _BULLET_START.match(text):
                        text = f"- {text}"
                    blocks.append(text)
            else:  # Table
                for row in item.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        blocks.append(" | ".join(cells))
        return "\n\n".join(blocks)

    raise ValueError(f"unsupported file type: {ext} (use .pdf, .docx, .txt or .md)")


# --------------------------------------------------------------------------
# chunking
# --------------------------------------------------------------------------

def _last_heading(block: str) -> str:
    """Most recent Markdown/underlined heading in a block, if any."""
    for line in reversed(block.splitlines()):
        s = line.strip()
        if s.startswith("#") or (s.endswith(":") and 3 < len(s) < 60 and not s.startswith(("-", "*"))):
            return s
    return ""


def chunk_text(text: str, size: int = 3500, overlap: int = 300) -> list[DocumentChunk]:
    """Paragraph-aware greedy packing.

    Splits on blank lines so requirement bullets are not cut in half, and
    carries the last-seen heading into the next chunk. Without the carryover a
    section that straddles a boundary loses its heading, and the extractor sees
    orphaned bullets with no idea whether they are in-scope items, out-of-scope
    items or constraints — which silently drops or mislabels requirements.
    """
    paragraphs = [p for p in text.split("\n\n") if p.strip()]
    chunks: list[DocumentChunk] = []
    buf: list[str] = []
    buf_len = 0
    cursor = 0

    carried = ""

    def flush() -> None:
        nonlocal buf, buf_len, cursor, carried
        if not buf:
            return
        body = "\n\n".join(buf)
        if carried and not body.lstrip().startswith("#"):
            body = f"{carried} (continued)\n\n{body}"
        carried = _last_heading(body) or carried
        chunks.append(DocumentChunk(
            chunk_id=f"C{len(chunks) + 1:03d}",
            text=body,
            order=len(chunks),
            char_start=cursor,
            char_end=cursor + len(body),
        ))
        cursor += len(body)
        tail = body[-overlap:] if overlap and len(body) > overlap else ""
        buf = [tail] if tail else []
        buf_len = len(tail)

    for para in paragraphs:
        if buf_len + len(para) > size and buf:
            flush()
        buf.append(para)
        buf_len += len(para) + 2
    flush()
    return chunks


# --------------------------------------------------------------------------
# summarization
# --------------------------------------------------------------------------

def summarize(
    text: str,
    client: LLMClient,
    *,
    title: str = "Business Requirements Summary",
    chunk_size: int = 3500,
    chunk_overlap: int = 300,
) -> BRSSummary:
    """Map over chunks, then reduce. One LLM call per chunk plus one reduce
    call; short documents skip the reduce step."""
    chunks = chunk_text(text, chunk_size, chunk_overlap)

    partials: list[dict] = []
    for i, ch in enumerate(chunks, start=1):
        data = client.complete_json(LLMRequest(
            task="summarize_chunk",
            system=SYSTEM_ANALYST,
            user=SUMMARIZE_CHUNK.format(i=i, n=len(chunks), text=ch.text),
            payload={"text": ch.text, "index": i, "total": len(chunks)},
        ))
        partials.append(data if isinstance(data, dict) else {})

    if len(partials) == 1:
        merged = dict(partials[0])
        merged["title"] = title
    else:
        import json
        merged = client.complete_json(LLMRequest(
            task="reduce_summary",
            system=SYSTEM_ANALYST,
            user=REDUCE_SUMMARY.format(partials=json.dumps(partials, ensure_ascii=False)[:60000]),
            payload={"partials": partials, "title": title},
        ))

    return BRSSummary(
        title=merged.get("title") or title,
        business_goals=list(merged.get("business_goals") or []),
        stakeholders=list(merged.get("stakeholders") or []),
        scope_in=list(merged.get("scope_in") or []),
        scope_out=list(merged.get("scope_out") or []),
        constraints=list(merged.get("constraints") or []),
        assumptions=list(merged.get("assumptions") or []),
        open_questions=list(merged.get("open_questions") or []),
        raw_char_count=len(text),
        chunk_count=len(chunks),
    )
