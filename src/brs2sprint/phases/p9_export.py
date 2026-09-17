"""Phase 9 — Export. JSON and Markdown use stdlib only; DOCX is optional."""

from __future__ import annotations

from pathlib import Path

from ..schemas import PipelineResult, SRS, to_json


def export_json(result: PipelineResult, out_dir: str) -> str:
    path = Path(out_dir) / f"{result.run_id}_full.json"
    path.write_text(to_json(result), encoding="utf-8")
    return str(path)


# --------------------------------------------------------------------------
# markdown
# --------------------------------------------------------------------------

def srs_markdown(srs: SRS) -> str:
    L: list[str] = [f"# {srs.title}", "", "## 1. Introduction", srs.introduction, "",
                    "## 2. Overall description", srs.overall_description, "",
                    "## 3. Requirements", ""]
    L += ["| ID | Type | Priority | Kind | Discipline | Requirement |",
          "|---|---|---|---|---|---|"]
    for r in srs.requirements:
        L.append(
            f"| {r.req_id} | {r.req_type.value} | {r.priority.value} | "
            f"{r.kind.value if r.kind else '-'} | "
            f"{r.discipline.value if r.discipline else '-'} | {r.text.replace('|', '/')} |"
        )
    L += ["", "### 3.1 Acceptance criteria", ""]
    for r in srs.requirements:
        L.append(f"**{r.req_id}** — {r.text}")
        L += [f"- {a}" for a in r.acceptance_criteria] or ["- _none provided_"]
        L.append("")
    if srs.use_cases:
        L += ["## 4. Use cases", ""]
        for u in srs.use_cases:
            L += [f"### {u.uc_id} — {u.name}", f"**Actor:** {u.actor}", "",
                  "**Preconditions**"] + [f"- {p}" for p in u.preconditions]
            L += ["", "**Main flow**"] + [f"{i}. {s}" for i, s in enumerate(u.main_flow, 1)]
            L += ["", "**Alternate flows**"] + [f"- {s}" for s in u.alt_flows] + [""]
    if srs.glossary:
        L += ["## 5. Glossary", ""] + [f"- **{k}** — {v}" for k, v in srs.glossary.items()]
    return "\n".join(L)


def full_markdown(result: PipelineResult) -> str:
    L: list[str] = [f"# Delivery pack — run `{result.run_id}`", "",
                    f"Source: `{result.source_path}`", ""]

    if result.summary:
        s = result.summary
        L += ["## Business summary", "", "**Goals**"] + [f"- {g}" for g in s.business_goals]
        L += ["", "**Stakeholders**"] + [f"- {x}" for x in s.stakeholders]
        L += ["", "**In scope**"] + [f"- {x}" for x in s.scope_in]
        L += ["", "**Out of scope**"] + [f"- {x}" for x in s.scope_out]
        L += ["", "**Constraints**"] + [f"- {x}" for x in s.constraints] + [""]

    if result.enriched and result.enriched.findings:
        L += ["## Identified gaps", ""]
        for f in result.enriched.findings:
            L += [f"- **{f.question}** — {f.answer} _(confidence {f.confidence})_"]
        L.append("")

    if result.srs:
        L += [srs_markdown(result.srs), ""]

    d = result.design
    if d and d.components:
        L += ["## High-level design", "", d.overview, "", "```mermaid", d.mermaid_hld, "```", "",
              "### Components", "", "| Component | Responsibility | Tech | Depends on |", "|---|---|---|---|"]
        L += [f"| {c.name} | {c.responsibility} | {c.tech} | {', '.join(c.depends_on) or '-'} |"
              for c in d.components]
        L += ["", "### Data flow", "", d.data_flow, ""]
        if d.api_contracts:
            L += ["### API contracts", "", "| Method | Path | Summary |", "|---|---|---|"]
            L += [f"| {a.method} | `{a.path}` | {a.summary} |" for a in d.api_contracts]
            L.append("")
        if d.db_tables:
            L += ["### Data model", "", "```mermaid", d.mermaid_erd or "", "```", ""]
            for t in d.db_tables:
                L.append(f"**{t.name}**")
                L += [f"- `{col}` — {typ}" for col, typ in t.columns.items()] + [""]

    if result.plan and result.plan.sprints:
        L += ["## Sprint plan", ""]
        for sp in result.plan.sprints:
            L += [f"### Sprint {sp.number} — {sp.committed_points}/{sp.capacity} points", "",
                  "| Ticket | Title | Points | Labels | Blocked by |", "|---|---|---|---|---|"]
            L += [f"| {t.ticket_id} | {t.title} | {t.story_points} | "
                  f"{', '.join(t.labels)} | {', '.join(t.depends_on) or '-'} |" for t in sp.tickets]
            L.append("")
        if result.plan.unscheduled:
            L += ["**Unscheduled**"] + [f"- {t.ticket_id} {t.title}" for t in result.plan.unscheduled]
        if result.plan.warnings:
            L += ["", "**Planner warnings**"] + [f"- {w}" for w in result.plan.warnings]
    return "\n".join(L)


def export_markdown(result: PipelineResult, out_dir: str) -> str:
    path = Path(out_dir) / f"{result.run_id}_delivery_pack.md"
    path.write_text(full_markdown(result), encoding="utf-8")
    return str(path)


# --------------------------------------------------------------------------
# docx (optional)
# --------------------------------------------------------------------------

def export_srs_docx(srs: SRS, out_dir: str, run_id: str) -> str:
    """Renders the SRS as a Word document. Needs `pip install python-docx`."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise ImportError("DOCX export needs `pip install python-docx`") from exc

    doc = Document()
    doc.add_heading(srs.title, level=0)
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(srs.introduction)
    doc.add_heading("2. Overall description", level=1)
    doc.add_paragraph(srs.overall_description)

    doc.add_heading("3. Requirements", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, ["ID", "Type", "Priority", "Kind", "Requirement"]):
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r in srs.requirements:
        row = table.add_row().cells
        row[0].text = r.req_id
        row[1].text = r.req_type.value
        row[2].text = r.priority.value
        row[3].text = r.kind.value if r.kind else "-"
        row[4].text = r.text

    doc.add_heading("3.1 Acceptance criteria", level=2)
    for r in srs.requirements:
        p = doc.add_paragraph()
        p.add_run(f"{r.req_id} — ").bold = True
        p.add_run(r.text)
        for a in r.acceptance_criteria:
            doc.add_paragraph(a, style="List Bullet")

    if srs.use_cases:
        doc.add_heading("4. Use cases", level=1)
        for u in srs.use_cases:
            doc.add_heading(f"{u.uc_id} — {u.name}", level=2)
            doc.add_paragraph(f"Actor: {u.actor}")
            for step in u.main_flow:
                doc.add_paragraph(step, style="List Number")

    for style_name in ("Normal",):
        doc.styles[style_name].font.size = Pt(10.5)

    path = Path(out_dir) / f"{run_id}_SRS.docx"
    doc.save(str(path))
    return str(path)
