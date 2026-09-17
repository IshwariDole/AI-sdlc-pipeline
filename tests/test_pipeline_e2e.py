"""End-to-end pipeline test using the offline mock provider.

Asserts the *contract* between phases (ids resolve, nothing is dropped, the
plan is feasible), not the wording of any LLM output — so it stays valid when
you swap in a real provider.
"""

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from brs2sprint.config import settings                         # noqa: E402
from brs2sprint.llm import MockClient                          # noqa: E402
from brs2sprint.phases import p1_ingest, p8_sprints            # noqa: E402
from brs2sprint.pipeline import PipelineOptions, run_pipeline  # noqa: E402

SAMPLE = ROOT / "data" / "samples" / "sample_brs.md"


def _run(**kw):
    settings.output_dir = tempfile.mkdtemp(prefix="b2s-test-")
    settings.db_path = str(Path(settings.output_dir) / "test.db")
    return run_pipeline(str(SAMPLE), client=MockClient(),
                        options=PipelineOptions(**kw), progress=lambda *a: None)


def test_full_run_produces_every_artefact():
    r = _run(velocity=25, max_sprints=30)
    assert r.summary and r.summary.business_goals
    assert r.srs and r.srs.requirements
    assert r.design and r.design.components
    assert r.tasks and r.tickets and r.plan
    assert r.llm_calls > 0


def test_requirement_ids_are_unique():
    r = _run(max_sprints=30)
    ids = [x.req_id for x in r.srs.requirements]
    assert len(ids) == len(set(ids))


def test_every_ticket_traces_back_to_a_requirement():
    r = _run(max_sprints=30)
    valid = {x.req_id for x in r.srs.requirements}
    for t in r.tickets:
        assert t.req_ids, f"{t.ticket_id} has no requirement link"
        assert set(t.req_ids) <= valid, f"{t.ticket_id} references unknown requirement"


def test_ticket_dependencies_all_resolve():
    r = _run(max_sprints=30)
    ids = {t.ticket_id for t in r.tickets}
    for t in r.tickets:
        for d in t.depends_on:
            assert d in ids, f"{t.ticket_id} depends on missing {d}"
            assert d != t.ticket_id


def test_plan_is_feasible():
    r = _run(velocity=30, max_sprints=40)
    where = {t.ticket_id: s.number for s in r.plan.sprints for t in s.tickets}
    for s in r.plan.sprints:
        assert s.committed_points <= s.capacity
        for t in s.tickets:
            for d in t.depends_on:
                if d in where:
                    assert where[d] < s.number, f"{t.ticket_id} scheduled before {d}"


def test_no_tickets_lost_between_phase_7_and_8():
    r = _run(velocity=30, max_sprints=40)
    placed = sum(len(s.tickets) for s in r.plan.sprints) + len(r.plan.unscheduled)
    assert placed == len(r.tickets)


def test_stop_after_halts_the_pipeline():
    r = _run(stop_after=3, max_sprints=30)
    assert r.srs is not None
    assert r.design is None and r.tickets == []


def test_chunking_preserves_all_paragraphs():
    text = SAMPLE.read_text(encoding="utf-8")
    chunks = p1_ingest.chunk_text(text, size=900, overlap=100)
    assert len(chunks) > 1
    joined = " ".join(c.text for c in chunks)
    for probe in ("no-show rates", "LabConnect", "procurement team", "data residency"):
        assert probe in joined, f"chunking dropped: {probe}"


def test_chunking_carries_headings_across_boundaries():
    text = SAMPLE.read_text(encoding="utf-8")
    chunks = p1_ingest.chunk_text(text, size=700, overlap=80)
    for c in chunks[1:]:
        head = c.text.lstrip().splitlines()[0]
        assert head.startswith("#") or "(continued)" in head, \
            "a continuation chunk lost its section heading"


def test_small_velocity_still_terminates():
    r = _run(velocity=8, max_sprints=5)
    m = p8_sprints.plan_metrics(r.plan)
    assert m["sprints"] <= 5
    assert r.plan.unscheduled          # expected: work remains


def test_docx_table_stays_under_its_own_heading():
    """Regression test: load_document used to append ALL tables after ALL
    paragraphs, silently detaching a requirements table from its section
    heading. This builds a minimal .docx with a heading, a bullet list, and
    a table — in that order — and asserts the loaded text preserves it."""
    import docx
    doc = docx.Document()
    doc.add_heading("1. Goals", level=1)
    doc.add_paragraph("Ship the thing.", style="List Bullet")
    doc.add_heading("2. Requirements", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Requirement"
    table.rows[1].cells[0].text = "BR-01"
    table.rows[1].cells[1].text = "The system shall log in a user."
    doc.add_heading("3. Constraints", level=1)
    doc.add_paragraph("Ship within budget.", style="List Bullet")

    tmp = Path(tempfile.mkdtemp()) / "mini.docx"
    doc.save(str(tmp))
    text = p1_ingest.load_document(tmp)

    req_pos = text.find("2. Requirements")
    table_pos = text.find("BR-01")
    constraints_pos = text.find("3. Constraints")
    assert req_pos != -1 and table_pos != -1 and constraints_pos != -1
    assert req_pos < table_pos < constraints_pos, \
        "table row landed outside its own section"


def test_docx_bullet_list_items_are_recoverable_as_bullets():
    """Word list paragraphs don't store the bullet glyph in .text — this
    checks load_document restores a marker so bullet-based extraction
    (the mock heuristics, or any regex-based downstream tool) still works."""
    import docx
    doc = docx.Document()
    doc.add_heading("Scope", level=1)
    doc.add_paragraph("Checkout with online payments", style="List Bullet")
    doc.add_paragraph("Order tracking", style="List Bullet")

    tmp = Path(tempfile.mkdtemp()) / "mini2.docx"
    doc.save(str(tmp))
    text = p1_ingest.load_document(tmp)

    from brs2sprint.llm.mock import _BULLET
    bullet_lines = [ln for ln in text.splitlines() if _BULLET.match(ln)]
    assert len(bullet_lines) == 2, \
        f"expected 2 recognisable bullets, got {len(bullet_lines)}: {text!r}"


if __name__ == "__main__":
    fns = [(k, v) for k, v in sorted(globals().items()) if k.startswith("test_")]
    failed = 0
    for name, fn in fns:
        try:
            fn()
            print(f"  PASS  {name}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL  {name}: {type(exc).__name__}: {exc}")
    print(f"\n{len(fns) - failed}/{len(fns)} passed")
    raise SystemExit(1 if failed else 0)
